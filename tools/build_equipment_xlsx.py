# 장비 DB 엑셀 생성: python3 tools/build_equipment_xlsx.py
# 입력: data/equipment/equipment_db.json (gen_equipment_db.py), balance_report.json, calibration.json
# 출력: docs/equipment_db.xlsx — 파란 글씨 = 입력값(수정 가능), 검정 = 수식, 초록 = 다른 시트 참조
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

ROOT = os.path.join(os.path.dirname(__file__), '..')
DB = json.load(open(os.path.join(ROOT, 'data/equipment/equipment_db.json'), encoding='utf-8'))
REP = json.load(open(os.path.join(ROOT, 'data/equipment/balance_report.json'), encoding='utf-8'))
CAL = json.load(open(os.path.join(ROOT, 'data/equipment/calibration.json'), encoding='utf-8'))

F = 'Arial'
f_in = Font(name=F, color='0000FF')
f_calc = Font(name=F, color='000000')
f_link = Font(name=F, color='008000')
f_head = Font(name=F, bold=True, color='FFFFFF')
f_title = Font(name=F, bold=True, size=14)
f_bold = Font(name=F, bold=True)
f_note = Font(name=F, italic=True, color='666666', size=9)
fill_head = PatternFill('solid', fgColor='3A3550')
fill_key = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='BBBBBB')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
PCT = '0.0%'
PCT2 = '0.00%'

GR = [g['code'] for g in DB['grades']]
STAT_NAME = {s['key']: s['name'] for s in DB['stats']}
STAT_PCT = {s['key']: s['pct'] for s in DB['stats']}
SLOT_NAME = {s['key']: s['name'] for s in DB['slots']}
CLS_NAME = {c['key']: c['name'] for c in DB['classes']}
CLS_NAME['common'] = '공통'

wb = Workbook()


def sheet(title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font = f_head; x.fill = fill_head; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); x.border = box
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[CL(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, r, c, v, font=None, fmt=None, bold=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = font or (f_calc if isinstance(v, str) and v.startswith('=') else f_in if isinstance(v, (int, float)) else Font(name=F, bold=bold))
    if fmt: x.number_format = fmt
    x.border = box
    return x


def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = f_title
    if sub: ws['A2'] = sub; ws['A2'].font = f_note


# ------------------------------------------------------------------ 0. 읽어보기
ws = sheet('읽어보기', first=True)
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 110
title(ws, '숲속 원정대 — 장비 DB (1단계)')
lines = [
    ('h', '규칙 (요청 사양)'),
    ('', '등급 8단계: UC < C < R < SR < SSR < UR < Legend(L) < Elder(E).'),
    ('', '부위: 직업별 무기 · 갑옷 · 소울 메달 / 공통 반지 · 목걸이 · 벨트.'),
    ('', '직업 무기·갑옷은 직업에 맞는 주 스탯 + 고유 스킬 효과가 붙는다. 추가 옵션은 R부터 1개씩 늘어 UR에서 최대 4개.'),
    ('', 'L: 보석 슬롯 1개 + (직업 특성 스탯 또는 스킬 보너스) 0~1개.'),
    ('', 'E: 보석 슬롯 1~3개 + 직업 특성 스탯 1개 + 스킬 보너스 0~2개.'),
    ('', '보석: 각 스테이지 보스 처치 보상. 등급별 효과 개수 R까지 1개 · UR까지 2개 · E까지 3개.'),
    ('h', '밸런스 원칙 (전략 RPG)'),
    ('', '큰 적은 [방어 태세](받는 피해 감소)와 [그로기 게이지]를 가진다. 기절·스킬·차지 끊기로 게이지를 깎으면 6초간 무방비(방어 해제 + 받는 피해 +50%).'),
    ('', '→ 비슷한 등급에서는 딜만 올려 밀어붙이는 플레이(딜찍누)보다 기믹 공략이 확실히 유리하도록 조정. 딜만으로 이기려면 2등급 이상 차이가 필요하도록 목표.'),
    ('', '→ 추가 옵션/보석에 공략 스탯(그로기 피해, 그로기 적 피해, 기절·도발 지속)을 넣어 장비로도 공략 빌드를 만들 수 있게 함.'),
    ('', '→ 순수 딜 옵션은 1줄당 전투력 약 1.1%, 생존/공략 옵션은 약 1.5%로 책정 (딜 옵션이 더 약함).'),
    ('h', '시트 안내'),
    ('', '등급 · 스탯 · 장비_기본 · 추가옵션 · 직업특성스탯 · 스킬보너스 · 보석 · 난이도 · 드랍테이블: 원본 데이터 (파란 글씨 = 수정 가능한 입력값).'),
    ('', '장비_등급별: 장비 49종 × 8등급의 실제 수치. 모두 수식(장비_기본 × 등급 배율) — 입력값을 바꾸면 자동 갱신.'),
    ('', '밸런스분석: 등급별 전투력, 옵션 공정성, 누적 상한, 전투 시뮬 승률(기믹 공략 vs 딜만), 딜찍누 지수. 시뮬 결과는 tools/equip_balance.cjs 실행 결과 값.'),
    ('h', '색 규칙'),
    ('', '파란 글씨 = 입력값 / 검정 = 수식 / 초록 = 다른 시트 참조 / 노란 칸 = 핵심 가정.'),
    ('h', '가정 · 확인이 필요한 점'),
    ('', '1) 등급 순서는 요청에 적힌 순서대로 UC를 최하위로 두었다 (일반적으로는 C < UC인 경우도 있음).'),
    ('', '2) L·E도 추가 옵션 4개를 유지 (UR 최대 4개 규칙의 연장). L·E의 차별점은 보석 슬롯 · 특성 스탯 · 스킬 보너스.'),
    ('', '3) 공통 장신구(반지·목걸이·벨트)의 L·E 직업 특성 스탯/스킬 보너스는 착용한 캐릭터의 직업 기준으로 붙는다.'),
    ('', '4) 소울 메달 = 직업 특성 스탯 주 스탯 + 필살기 충전. 무기·갑옷과 겹치지 않게 고유 효과는 메달 전용 라인 2종.'),
    ('', '5) 장비가 원정 사이에 유지(영구 성장)된다고 가정 → 적 강함을 8단계 난이도로 나눔 (난이도 시트). 데모의 기본 던전 = 견습(T1) 아래 단계.'),
    ('', '6) "각 스테이지 보스" — 현재 데모는 5스테이지 보스 1종뿐. 보석 드랍은 보스 노드 기준으로 정의 (스테이지별 중간 보스 추가 시 동일 테이블 사용).'),
    ('', '7) 승률 수치는 헤드리스 시뮬(3개 파티 조합 × 시드, 4스테이지 일반/정예 + 보스) 기준이며 실제 플레이어 실력에 따라 달라진다.'),
]
r = 3
for kind, t in lines:
    c = ws.cell(row=r, column=1 if kind == 'h' else 2, value=t)
    c.font = Font(name=F, bold=True, size=11) if kind == 'h' else Font(name=F)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    r += 1 if kind else 1
    if kind == 'h': r += 0

# ------------------------------------------------------------------ 1. 등급
ws = sheet('등급')
title(ws, '등급', '주 스탯 배율 · 옵션 개수 · 옵션 배율 · 고유 효과 배율 · 보석 슬롯 · 특성/스킬 보너스 · 보석 등급 효과')
cols = ['등급', '이름', '주 스탯 배율', '추가 옵션 수', '옵션 수치 배율', '고유 효과 배율', '보석 슬롯 최소', '보석 슬롯 최대', '특성 스탯 최소', '특성 스탯 최대', '스킬 보너스 최소', '스킬 보너스 최대', '보석 효과 개수', '보석 수치 배율', '주 스탯 증가율(전 등급 대비)']
header(ws, 4, cols, [8, 10, 12, 11, 12, 12, 11, 11, 11, 11, 12, 12, 11, 12, 16])
for i, g in enumerate(DB['grades']):
    rr = 5 + i
    vals = [g['code'], g['name'], g['main'], g['opt'], g['optMult'], g['innate'], g['gem'][0], g['gem'][1], g['cls'][0], g['cls'][1], g['skill'][0], g['skill'][1], DB['gemEffectCount'][g['code']], DB['gemMult'][g['code']]]
    for j, v in enumerate(vals, 1):
        put(ws, rr, j, v, font=Font(name=F, bold=True) if j <= 2 else None, fmt='0.00' if j in (3, 5, 6, 14) else None)
    put(ws, rr, 15, f'=IF(ROW()=5,"-",C{rr}/C{rr - 1}-1)', fmt=PCT)
GROW = {g: 5 + i for i, g in enumerate(GR)}
GRANGE = '등급!$A$5:$A$12'
def gref(col, grade_cell):
    return f'INDEX(등급!${col}$5:${col}$12,MATCH({grade_cell},{GRANGE},0))'
r = 14
put(ws, r, 1, 'L 추가 효과 확률', font=f_bold); r += 1
for k, v in DB['lExtraWeights'].items():
    put(ws, r, 1, {'none': '없음', 'cls': '특성 스탯', 'skill': '스킬 보너스'}[k], font=Font(name=F)); put(ws, r, 2, v); put(ws, r, 3, f'=B{r}/SUM($B$15:$B$17)', fmt=PCT); r += 1
r += 1
put(ws, r, 1, 'E 보석 슬롯 확률', font=f_bold); r += 1; s0 = r
for k, v in DB['eGemWeights'].items():
    put(ws, r, 1, f'{k}칸', font=Font(name=F)); put(ws, r, 2, v); r += 1
for rr in range(s0, r): put(ws, rr, 3, f'=B{rr}/SUM($B${s0}:$B${r - 1})', fmt=PCT)
r += 1
put(ws, r, 1, 'E 스킬 보너스 확률', font=f_bold); r += 1; s0 = r
for k, v in DB['eSkillWeights'].items():
    put(ws, r, 1, f'{k}개', font=Font(name=F)); put(ws, r, 2, v); r += 1
for rr in range(s0, r): put(ws, rr, 3, f'=B{rr}/SUM($B${s0}:$B${r - 1})', fmt=PCT)
ws.cell(row=r + 1, column=1, value='가중치는 상대값. 확률 = 가중치 ÷ 합계.').font = f_note

# ------------------------------------------------------------------ 2. 스탯
ws = sheet('스탯')
title(ws, '스탯', '전투력 가중치(pw): 스탯 1%p가 전투력 몇 %에 해당하는지 (밸런스 분석용 근사)')
header(ws, 4, ['키', '이름', '단위', '전투력 가중치', '상한', '공략 스탯', '비고'], [12, 16, 8, 14, 10, 10, 50])
for i, s in enumerate(DB['stats']):
    rr = 5 + i
    put(ws, rr, 1, s['key'], font=Font(name=F)); put(ws, rr, 2, s['name'], font=Font(name=F)); put(ws, rr, 3, s['unit'] or '수치', font=Font(name=F))
    put(ws, rr, 4, s['pw'] if s['pw'] is not None else '환산', fmt='0.00') if s['pw'] is not None else put(ws, rr, 4, '환산', font=Font(name=F))
    if s.get('cap') is not None: put(ws, rr, 5, s['cap'], fmt='0%')
    else: put(ws, rr, 5, '-', font=Font(name=F))
    put(ws, rr, 6, 'O' if s.get('gimmick') else '', font=Font(name=F))
    put(ws, rr, 7, s.get('note', ''), font=Font(name=F))
STAT_ROWS = len(DB['stats'])

# ------------------------------------------------------------------ 3. 장비_기본
ws = sheet('장비_기본')
title(ws, '장비 기본값 (UC 기준)', '등급별 수치 = 기본값 × 등급 배율. 공격력·체력은 수치, 나머지는 비율.')
cols = ['ID', '이름', '부위', '직업', '라인', '주 스탯1', '기본값1', '주 스탯2', '기본값2', '고유 효과', '고유 기본값', '단위', '고유 효과 키']
header(ws, 4, cols, [18, 20, 10, 8, 6, 14, 10, 14, 10, 40, 11, 7, 16])
ITEM_ROW = {}
for i, it in enumerate(DB['items']):
    rr = 5 + i
    ITEM_ROW[it['id']] = rr
    m = it['main'] + [None] * (2 - len(it['main']))
    put(ws, rr, 1, it['id'], font=Font(name=F)); put(ws, rr, 2, it['name'], font=Font(name=F, bold=True))
    put(ws, rr, 3, SLOT_NAME[it['slot']], font=Font(name=F)); put(ws, rr, 4, CLS_NAME[it['cls']], font=Font(name=F)); put(ws, rr, 5, it['line'])
    for k, mm in enumerate(m):
        c = 6 + k * 2
        if mm:
            put(ws, rr, c, STAT_NAME[mm['stat']], font=Font(name=F))
            put(ws, rr, c + 1, mm['base'], fmt='0.0' if not STAT_PCT[mm['stat']] else PCT2)
        else:
            put(ws, rr, c, '-', font=Font(name=F)); put(ws, rr, c + 1, '', font=Font(name=F))
    inn = it['innate']
    if inn:
        put(ws, rr, 10, inn['text'].replace('{v}', '[값]'), font=Font(name=F)); put(ws, rr, 11, inn['base'], fmt='0.0'); put(ws, rr, 12, inn['unit'], font=Font(name=F)); put(ws, rr, 13, inn['key'], font=Font(name=F))
    else:
        for c in (10, 11, 12, 13): put(ws, rr, c, '-' if c == 10 else '', font=Font(name=F))
N_ITEMS = len(DB['items'])

# ------------------------------------------------------------------ 4. 장비_등급별 (수식)
ws = sheet('장비_등급별')
title(ws, '장비 × 등급 수치 (수식)', '주 스탯 = 기본값 × 주 스탯 배율 / 고유 효과 = 고유 기본값 × 고유 효과 배율. 옵션·보석 슬롯은 등급 시트 참조.')
cols = ['이름', '부위', '직업', '등급', '주 스탯1', '값1', '주 스탯2', '값2', '고유 효과', '고유 값', '추가 옵션 수', '보석 슬롯', '특성 스탯', '스킬 보너스']
header(ws, 4, cols, [20, 9, 7, 7, 14, 10, 14, 10, 40, 9, 10, 10, 10, 11])
rr = 5
for it in DB['items']:
    br = ITEM_ROW[it['id']]
    for g in GR:
        put(ws, rr, 1, f'=장비_기본!B{br}', font=f_link); put(ws, rr, 2, f'=장비_기본!C{br}', font=f_link); put(ws, rr, 3, f'=장비_기본!D{br}', font=f_link)
        put(ws, rr, 4, g, font=Font(name=F, bold=True))
        put(ws, rr, 5, f'=장비_기본!F{br}', font=f_link)
        f1 = '0.0' if not STAT_PCT[it['main'][0]['stat']] else PCT2
        put(ws, rr, 6, f'=장비_기본!G{br}*{gref("C", f"$D{rr}")}', fmt=f1)
        if len(it['main']) > 1:
            put(ws, rr, 7, f'=장비_기본!H{br}', font=f_link)
            f2 = '0.0' if not STAT_PCT[it['main'][1]['stat']] else PCT2
            put(ws, rr, 8, f'=장비_기본!I{br}*{gref("C", f"$D{rr}")}', fmt=f2)
        else:
            put(ws, rr, 7, '-', font=Font(name=F)); put(ws, rr, 8, '', font=Font(name=F))
        if it['innate']:
            put(ws, rr, 9, f'=장비_기본!J{br}', font=f_link)
            put(ws, rr, 10, f'=ROUND(장비_기본!K{br}*{gref("F", f"$D{rr}")},1)', fmt='0.0')
        else:
            put(ws, rr, 9, '-', font=Font(name=F)); put(ws, rr, 10, '', font=Font(name=F))
        put(ws, rr, 11, f'={gref("D", f"$D{rr}")}')
        put(ws, rr, 12, f'=IF({gref("H", f"$D{rr}")}=0,"-",IF({gref("G", f"$D{rr}")}={gref("H", f"$D{rr}")},{gref("G", f"$D{rr}")}&"칸",{gref("G", f"$D{rr}")}&"~"&{gref("H", f"$D{rr}")}&"칸"))')
        put(ws, rr, 13, f'=IF({gref("J", f"$D{rr}")}=0,"-",IF({gref("I", f"$D{rr}")}={gref("J", f"$D{rr}")},{gref("I", f"$D{rr}")}&"개",{gref("I", f"$D{rr}")}&"~"&{gref("J", f"$D{rr}")}&"개"))')
        put(ws, rr, 14, f'=IF({gref("L", f"$D{rr}")}=0,"-",{gref("K", f"$D{rr}")}&"~"&{gref("L", f"$D{rr}")}&"개")')
        for c in (11, 12, 13, 14): ws.cell(row=rr, column=c).font = f_link
        rr += 1
ws.auto_filter.ref = f'A4:N{rr - 1}'

# ------------------------------------------------------------------ 5. 추가옵션
ws = sheet('추가옵션')
title(ws, '추가 옵션 풀', 'R 기준 범위 × 등급 옵션 배율. 한 장비에 같은 옵션은 1줄만. 등장 확률 = 가중치 ÷ 해당 부위 합계.')
cols = ['키', '옵션', 'R 최소', 'R 최대', '가중치', '붙는 부위'] + [f'{g} 최소' for g in GR[2:]] + [f'{g} 최대' for g in GR[2:]] + ['1줄 전투력(%)']
header(ws, 4, cols, [11, 16, 9, 9, 8, 30] + [9] * 12 + [12])
skey_row = {s['key']: 5 + i for i, s in enumerate(DB['stats'])}
for i, o in enumerate(DB['options']):
    rr = 5 + i
    put(ws, rr, 1, o['key'], font=Font(name=F)); put(ws, rr, 2, STAT_NAME[o['key']], font=Font(name=F, bold=True))
    put(ws, rr, 3, o['min'], fmt=PCT); put(ws, rr, 4, o['max'], fmt=PCT); put(ws, rr, 5, o['weight'])
    put(ws, rr, 6, '전 부위' if o['slots'] == '*' else ' · '.join(SLOT_NAME[s] for s in o['slots'].split(',')), font=Font(name=F))
    for k, g in enumerate(GR[2:]):
        put(ws, rr, 7 + k, f'=$C{rr}*등급!$E${GROW[g]}', fmt=PCT, font=f_link)
        put(ws, rr, 13 + k, f'=$D{rr}*등급!$E${GROW[g]}', fmt=PCT, font=f_link)
    put(ws, rr, 19, f'=(C{rr}+D{rr})/2*100*스탯!$D${skey_row[o["key"]]}', fmt='0.00')
nopt = len(DB['options'])
ws.cell(row=6 + nopt, column=1, value='1줄 전투력 = R 평균값(%p) × 전투력 가중치. 딜 옵션(공격력% 등) ≈ 1.1, 생존·공략 옵션 ≈ 1.5 — 딜찍누 억제를 위해 의도적으로 딜 옵션을 낮게 책정.').font = f_note

# ------------------------------------------------------------------ 6. 직업특성스탯
ws = sheet('직업특성스탯')
title(ws, '직업 특성 스탯 (L·E 전용)', 'L: 0~1개(스킬 보너스와 택1), E: 1개. E = L (상한 초과 누적 방지를 위해 동일 값)')
header(ws, 4, ['직업', '캐릭터', '특성 스탯', 'L 값', 'E 값', '기본 공격력', '기본 체력'], [10, 10, 16, 10, 10, 12, 12])
for i, c in enumerate(DB['classes']):
    rr = 5 + i
    cs = DB['classStat'][c['key']]
    put(ws, rr, 1, c['name'], font=Font(name=F, bold=True)); put(ws, rr, 2, c['heroName'], font=Font(name=F)); put(ws, rr, 3, STAT_NAME[cs['stat']], font=Font(name=F))
    put(ws, rr, 4, cs['L'], fmt=PCT); put(ws, rr, 5, f'=D{rr}', fmt=PCT); put(ws, rr, 6, c['baseAtk']); put(ws, rr, 7, c['baseHp'])
ws['E5'].comment = Comment('E 특성 스탯 = L 값. E가 3부위 이상 겹치면 특성 스탯이 상한을 넘기 때문 (밸런스 분석 결과).', 'balance')

# ------------------------------------------------------------------ 7. 스킬보너스
ws = sheet('스킬보너스')
title(ws, '스킬 보너스 (L·E 전용)', 'E 값 = ROUND(L × E 배율). 단, 횟수·초 단위 효과는 직접 지정(파란 글씨). 직업 전용 + 공통 풀에서 뽑힘.')
put(ws, 3, 1, 'E 배율', font=f_bold); x = put(ws, 3, 2, DB['skillBonusERatio'], fmt='0.00'); x.fill = fill_key
header(ws, 4, ['직업', '키', '효과', 'L 값', 'E 값', '단위', 'E/L'], [10, 12, 40, 9, 9, 7, 8])
rr = 5
for ck, lst in DB['skillBonus'].items():
    for b in lst:
        put(ws, rr, 1, CLS_NAME[ck], font=Font(name=F)); put(ws, rr, 2, b['key'], font=Font(name=F)); put(ws, rr, 3, b['text'].replace('{v}', '[값]'), font=Font(name=F))
        put(ws, rr, 4, b['L'])
        if round(b['L'] * DB['skillBonusERatio']) == b['E']: put(ws, rr, 5, f'=ROUND(D{rr}*$B$3,0)')
        else: put(ws, rr, 5, b['E'])  # 정수/초 단위라 배율 대신 직접 지정 (파란 글씨)
        put(ws, rr, 6, b['unit'], font=Font(name=F)); put(ws, rr, 7, f'=E{rr}/D{rr}', fmt='0.00')
        rr += 1
SB_END = rr - 1

# ------------------------------------------------------------------ 8. 보석
ws = sheet('보석')
title(ws, '보석 (스테이지 보스 보상)', '주 효과 = 기본값 × 보석 수치 배율. 2·3번째 효과 = 다른 보석 효과 또는 스킬 효과, 값 × 추가 효과 비율.')
put(ws, 3, 1, '추가 효과 비율', font=f_bold); x = put(ws, 3, 2, DB['gemExtraRatio'], fmt='0.00'); x.fill = fill_key
cols = ['키', '보석', '효과 스탯', '기본값'] + GR
header(ws, 4, cols, [10, 10, 16, 9] + [9] * 8)
for i, gm in enumerate(DB['gems']):
    rr = 5 + i
    put(ws, rr, 1, gm['key'], font=Font(name=F)); put(ws, rr, 2, gm['name'], font=Font(name=F, bold=True)); put(ws, rr, 3, STAT_NAME[gm['stat']], font=Font(name=F)); put(ws, rr, 4, gm['base'], fmt=PCT2)
    for k, g in enumerate(GR): put(ws, rr, 5 + k, f'=$D{rr}*등급!$N${GROW[g]}', fmt=PCT2, font=f_link)
r0 = 5 + len(DB['gems']) + 1
put(ws, r0, 1, '스킬 효과(추가 효과 전용)', font=f_bold)
header_r = r0 + 1
for j, c in enumerate(['키', '효과', '', '기본값'] + GR, 1):
    x = ws.cell(row=header_r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
for i, ex in enumerate(DB['gemSkillExtras']):
    rr = header_r + 1 + i
    put(ws, rr, 1, ex['key'], font=Font(name=F)); put(ws, rr, 2, ex['name'], font=Font(name=F)); put(ws, rr, 3, '', font=Font(name=F)); put(ws, rr, 4, ex['base'], fmt=PCT2)
    for k, g in enumerate(GR): put(ws, rr, 5 + k, f'=$D{rr}*등급!$N${GROW[g]}*$B$3', fmt=PCT2, font=f_link)
rr = header_r + 2 + len(DB['gemSkillExtras'])
put(ws, rr, 1, '효과 개수', font=f_bold)
for k, g in enumerate(GR): put(ws, rr, 5 + k, f'=등급!$M${GROW[g]}', font=f_link)

# ------------------------------------------------------------------ 9. 난이도
ws = sheet('난이도')
title(ws, '난이도 단계 (영구 장비 성장 대응)', '적 배율 = 권장 등급 풀세트로 기믹 공략 시 승률 85%가 되는 지점 (시뮬 보정 → 로그 회귀)')
header(ws, 4, ['단계', '이름', '적 체력 배율', '적 공격력 배율', '권장 장비', '보정 실측(85% 지점)', '회귀 곡선', '전 단계 대비'], [7, 10, 12, 13, 10, 18, 11, 12])
for i, t in enumerate(DB['tiers']):
    rr = 5 + i
    put(ws, rr, 1, t['tier']); put(ws, rr, 2, t['name'], font=Font(name=F, bold=True)); put(ws, rr, 3, t['hp'], fmt='0.00'); put(ws, rr, 4, f'=C{rr}', fmt='0.00'); put(ws, rr, 5, t['recGrade'], font=Font(name=F))
    put(ws, rr, 6, CAL['calib'][t['recGrade']], fmt='0.000'); put(ws, rr, 7, CAL['calibFit']['fitted'][t['recGrade']], fmt='0.000')
    put(ws, rr, 8, '-' if i == 0 else f'=C{rr}/C{rr - 1}-1', fmt=PCT, font=Font(name=F) if i == 0 else None)
ws['F4'].comment = Comment('tools/equip_balance.cjs --calibrate (시드 5개/조합): 적 배율을 이분 탐색해 4스테이지 일반·정예·보스 평균 승률이 85%가 되는 값. 시드 수가 적어 잡음이 있어 G열 회귀 곡선을 채택.', 'balance')
ws.cell(row=14, column=1, value='배율 1.00 = 데모 기본 던전(장비 없음). 견습(T1)부터 장비를 갖춘 원정대를 가정.').font = f_note

# ------------------------------------------------------------------ 10. 드랍테이블
ws = sheet('드랍테이블')
title(ws, '드랍 테이블 (등급 확률)', '가중치(입력) → 확률(수식). 출처별 중심 등급: 일반 -0.6, 정예 +0.4, 보스 +1.0, 상점 -1.0 (난이도 단계 기준).')
src_name = {s['key']: s['name'] for s in DB['dropSources']}
r = 4
for sk, tbl in DB['drops'].items():
    put(ws, r, 1, src_name[sk], font=f_bold); r += 1
    header_cols = ['난이도'] + [f'{g} 가중치' for g in GR] + [f'{g} 확률' for g in GR]
    for j, c in enumerate(header_cols, 1):
        x = ws.cell(row=r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box; x.alignment = Alignment(horizontal='center', wrap_text=True)
    r += 1
    for t, row in tbl.items():
        put(ws, r, 1, f'T{t}', font=Font(name=F))
        for k, v in enumerate(row): put(ws, r, 2 + k, v, fmt='0.00')
        for k in range(8): put(ws, r, 10 + k, f'=IF(SUM($B{r}:$I{r})=0,0,{CL(2 + k)}{r}/SUM($B{r}:$I{r}))', fmt=PCT)
        r += 1
    r += 1
ws.column_dimensions['A'].width = 26
for c in range(2, 18): ws.column_dimensions[CL(c)].width = 9

# ------------------------------------------------------------------ 11. 밸런스분석
ws = sheet('밸런스분석')
title(ws, '밸런스 분석', '전투력(%) = 장비 풀세트(6부위)가 올려주는 전투력 근사. 승률 = 헤드리스 전투 시뮬.')
ws.column_dimensions['A'].width = 16
for c in range(2, 20): ws.column_dimensions[CL(c)].width = 11
r = 4
put(ws, r, 1, '① 등급별 풀세트 전투력(%) — 직업별 · 기대값', font=f_bold); r += 1
for j, c in enumerate(['직업'] + GR + ['E/UR'], 1):
    x = ws.cell(row=r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
r += 1; p0 = r
for ck in REP['power']:
    put(ws, r, 1, CLS_NAME[ck], font=Font(name=F, bold=True))
    for k, g in enumerate(GR): put(ws, r, 2 + k, REP['power'][ck][g]['total'], fmt='0.0')
    put(ws, r, 10, f'=I{r}/G{r}', fmt='0.00')
    r += 1
put(ws, r, 1, '평균', font=f_bold)
for k in range(9): put(ws, r, 2 + k, f'=AVERAGE({CL(2 + k)}{p0}:{CL(2 + k)}{r - 1})', fmt='0.0' if k < 8 else '0.00')
put(ws, r + 1, 1, '직업 편차', font=f_bold)
for k in range(8): put(ws, r + 1, 2 + k, f'=MAX({CL(2 + k)}{p0}:{CL(2 + k)}{r - 1})/MIN({CL(2 + k)}{p0}:{CL(2 + k)}{r - 1})-1', fmt=PCT)
ws.cell(row=r + 2, column=1, value='직업 간 편차가 작을수록 공정. 목표 ±5% 이내.').font = f_note
r += 4

put(ws, r, 1, '② 전투력 구성 (탱커 기준, %)', font=f_bold); r += 1
parts = ['main', 'innate', 'options', 'classStat', 'skillBonus', 'gems']
pnames = ['주 스탯', '고유 효과', '추가 옵션', '특성 스탯', '스킬 보너스', '보석']
for j, c in enumerate(['구성'] + GR, 1):
    x = ws.cell(row=r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
r += 1
for pk, pn in zip(parts, pnames):
    put(ws, r, 1, pn, font=Font(name=F))
    for k, g in enumerate(GR): put(ws, r, 2 + k, REP['power']['tank'][g][pk], fmt='0.0')
    r += 1
r += 1

put(ws, r, 1, '③ 최대 누적 vs 상한 (E 풀세트 최악의 경우)', font=f_bold); r += 1
for j, c in enumerate(['스탯', '최대 누적', '상한', '상한 초과', '출처'], 1):
    x = ws.cell(row=r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
r += 1
for s in REP['stackMax']:
    put(ws, r, 1, s['name'], font=Font(name=F)); put(ws, r, 2, s['max'], fmt=PCT)
    if s.get('cap') is not None:
        put(ws, r, 3, s['cap'], fmt=PCT); put(ws, r, 4, f'=IF(B{r}>C{r},"초과 → 상한 적용","OK")')
    else:
        put(ws, r, 3, '-', font=Font(name=F)); put(ws, r, 4, '상한 없음', font=Font(name=F))
    put(ws, r, 5, s.get('sources', ''), font=Font(name=F))
    r += 1
ws.cell(row=r, column=1, value='초과분은 게임에서 상한으로 잘림 (CONST.STAT_CAPS). 일부러 한 스탯을 몰아도 상한까지만 → 빌드 다양성 유도.').font = f_note
r += 2

cells = REP['sim']['cells']
fights = REP['sim']['fights']
fname = {'stage4': '4스테이지 일반', 'elite4': '4스테이지 정예', 'boss': '보스'}
for prof, pname in (('gimmick', '④ 시뮬 승률 — 기믹 공략 (기절로 차지 끊기, 그로기에 필살기 집중, 위험 범위 회피)'), ('brute', '⑤ 시뮬 승률 — 딜만 (쿨마다 스킬 난사, 가까운 적만 공격)')):
    put(ws, r, 1, pname, font=f_bold); r += 1
    ws.cell(row=r, column=1, value='행 = 난이도, 열 = 장비 등급. 값 = 일반/정예/보스 평균 승률. 대각선(권장 등급)이 기믹 기준 85% 근처가 목표.').font = f_note; r += 1
    gcols = ['없음'] + GR
    for j, c in enumerate(['난이도'] + gcols, 1):
        x = ws.cell(row=r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
    r += 1
    for t in DB['tiers']:
        put(ws, r, 1, f"T{t['tier']} {t['name']} ({t['recGrade']})", font=Font(name=F))
        for k, g in enumerate(gcols):
            c = next((c for c in cells if c['tier'] == t['tier'] and c['grade'] == g and c['profile'] == prof), None)
            if c:
                v = sum(c[f]['win'] for f in fights) / len(fights)
                x = put(ws, r, 2 + k, round(v, 3), fmt='0%')
                x.font = f_calc
                if g == t['recGrade']: x.fill = PatternFill('solid', fgColor='FFF2B3')
                if g == t['recGrade'] or (prof == 'brute'): pass
        r += 1
    r += 1

put(ws, r, 1, '⑥ 딜찍누 지수 — 같은 등급에서 기믹 vs 딜만, 딜만으로 85%를 넘으려면 필요한 등급 차이', font=f_bold); r += 1
for j, c in enumerate(['난이도', '권장 등급', '기믹 승률', '딜만 승률', '차이', '딜만 85%에 필요한 추가 등급'], 1):
    x = ws.cell(row=r, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
r += 1
for b in REP['bruteIndex']:
    put(ws, r, 1, f"T{b['tier']}"); ws.cell(row=r, column=1).font = Font(name=F)
    put(ws, r, 2, b['recGrade'], font=Font(name=F)); x = put(ws, r, 3, b['gimmickWin'], fmt='0%'); x.font = f_calc
    x = put(ws, r, 4, b['bruteWin'], fmt='0%'); x.font = f_calc
    put(ws, r, 5, f'=C{r}-D{r}', fmt='+0%;-0%;0%')
    nb = b['bruteNeedsGradesAbove']
    put(ws, r, 6, '장비 범위 안에서 불가' if nb is None else f'+{nb}등급', font=Font(name=F))
    r += 1
ws.cell(row=r, column=1, value='시뮬 결과 값은 tools/equip_balance.cjs (SEEDS=4, 조합 3개) 실행 결과. 재생성: node tools/equip_balance.cjs → python3 tools/build_equipment_xlsx.py').font = f_note

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None, size=c.font.size if c.font else 11, italic=c.font.italic if c.font else False)

out = os.path.join(ROOT, 'docs/equipment_db.xlsx')
wb.save(out)
print('saved', out)
